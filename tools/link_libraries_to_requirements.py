#!/usr/bin/env python3
"""将 8.19 题库逐行关联到《急诊与麻醉科缺失.xlsx》的权威属性。"""
from collections import Counter, defaultdict
from copy import copy
from pathlib import Path
import json

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
LIBRARY_DIR = ROOT / "8.19"
REQUIREMENTS_PATH = ROOT / "data" / "requirements" / "急诊与麻醉科缺失.xlsx"
OLD_RESULT_FILES = {
    "A1": LIBRARY_DIR / "87676221-c539-494d-9dbc-b7cb706b22b3_1787166045318_success.jsonl",
    "A2": LIBRARY_DIR / "b980cead-a128-4d36-b302-8964e3ff0808_1787173956550_success.jsonl",
    "A3": LIBRARY_DIR / "8f8f2a6a-5432-4f56-ad15-adcdb1939517_1787176322622_success.jsonl",
    "A4": LIBRARY_DIR / "abfb4702-5736-43a1-a4f4-754ee45e28db_1787175996825_success.jsonl",
    "病例分析题": LIBRARY_DIR / "1ef21092-bd98-42fc-a964-dd67cafbeb2d_1787181086323_success.jsonl",
}


def load_jsonl(path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line]


def short_base(base):
    return base.replace("基地", "").replace("科", "")


def load_requirements():
    ws = openpyxl.load_workbook(REQUIREMENTS_PATH, read_only=True, data_only=True)["Sheet1"]
    by_type = defaultdict(list)
    for base, department, disease, a1, a2, a34, case in list(ws.values)[1:]:
        for question_type, enabled in (
            ("A1", a1),
            ("A2", a2),
            ("A3", a34),
            ("A4", a34),
            ("病例分析题", case),
        ):
            if enabled:
                by_type[question_type].append((base, department, disease))
    return by_type


def resolve_requirement(question_type, custom_id, requirements):
    parts = custom_id.rsplit("-", 1)[0].split("-")[1:]
    if question_type == "A1":
        disease = "-".join(parts)
        candidates = [item for item in requirements[question_type] if item[2] == disease]
    else:
        base_short, disease = parts[0], "-".join(parts[1:])
        candidates = [
            item
            for item in requirements[question_type]
            if short_base(item[0]) == base_short and item[2] == disease
        ]
    if len(candidates) != 1:
        raise ValueError(f"{question_type} / {custom_id} 未能唯一匹配要求文件：{candidates}")
    return candidates[0]


def response_for_source(source, old_responses, new_response_cache):
    if isinstance(source, int):
        if not 1 <= source <= len(old_responses):
            raise ValueError(f"旧批次来源行超出范围：{source}")
        return old_responses[source - 1]
    if isinstance(source, str) and source.startswith("820/"):
        file_name, line_text = source.split(":", 1)
        line_number = int(line_text.split()[0])
        path = ROOT / file_name
        responses = new_response_cache.get(path)
        if responses is None:
            responses = load_jsonl(path)
            new_response_cache[path] = responses
        if not 1 <= line_number <= len(responses):
            raise ValueError(f"820 来源行超出范围：{source}")
        return responses[line_number - 1]
    raise ValueError(f"未知来源行格式：{source!r}")


def copy_header_style(ws, column):
    source = ws.cell(1, 1)
    target = ws.cell(1, column)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def update_library(question_type, requirements, old_responses):
    path = LIBRARY_DIR / f"{question_type}.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    if headers[2:5] != ["专业基地", "科室", "病种"]:
        ws.insert_cols(3, 3)
        for column, header in zip((3, 4, 5), ("专业基地", "科室", "病种")):
            copy_header_style(ws, column)
            ws.cell(1, column).value = header

    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    ai_column = headers.index("考核要点") + 1 if "考核要点" in headers else headers.index("AI生成考核要点（参考）") + 1
    ws.cell(1, ai_column).value = "AI生成考核要点（参考）"
    requirement_column = ai_column + 1
    if ws.cell(1, requirement_column).value != "考核要点（要求文件）":
        copy_header_style(ws, requirement_column)
        ws.cell(1, requirement_column).value = "考核要点（要求文件）"

    new_response_cache = {}
    counts = Counter()
    unresolved = []
    for row in range(2, ws.max_row + 1):
        try:
            response = response_for_source(ws.cell(row, 2).value, old_responses, new_response_cache)
            base, department, disease = resolve_requirement(question_type, response["custom_id"], requirements)
        except Exception as exc:
            unresolved.append(f"第{row}行：{exc}")
            continue
        ws.cell(row, 3).value = base
        ws.cell(row, 4).value = department
        ws.cell(row, 5).value = disease
        ws.cell(row, requirement_column).value = disease
        counts[(base, department, disease)] += 1

    if unresolved:
        raise RuntimeError(f"{question_type} 存在未匹配记录：{unresolved[:5]}")

    for column, width in ((3, 18), (4, 30), (5, 40), (ai_column, 34), (requirement_column, 32)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    wb.save(path)
    return counts, ws.max_row - 1


def main():
    requirements = load_requirements()
    old_responses = {question_type: load_jsonl(path) for question_type, path in OLD_RESULT_FILES.items()}
    report = {}
    for question_type in OLD_RESULT_FILES:
        counts, row_count = update_library(question_type, requirements, old_responses[question_type])
        report[question_type] = {
            "row_count": row_count,
            "requirement_combinations": [
                {"专业基地": base, "科室": department, "病种": disease, "题目行数": count}
                for (base, department, disease), count in sorted(counts.items())
            ],
        }
    report_path = LIBRARY_DIR / "requirement_mapping_report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
