#!/usr/bin/env python3
"""删除题库中的 AI 生成考核要点列，仅保留要求文件考核要点。"""
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]


def main():
    for name in ("A1.xlsx", "A2.xlsx", "A3.xlsx", "A4.xlsx", "病例分析题.xlsx"):
        path = ROOT / "8.19" / name
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]
        if "AI生成考核要点（参考）" not in headers:
            raise RuntimeError(f"{name} 缺少 AI 生成考核要点列")
        ws.delete_cols(headers.index("AI生成考核要点（参考）") + 1)
        wb.save(path)
        print(f"{name}: 已删除 AI生成考核要点（参考）列")


if __name__ == "__main__":
    main()
