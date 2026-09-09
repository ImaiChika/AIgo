#!/usr/bin/env python3
"""清空旧病例分析题并按不定项病例串格式导入 824 试跑结果。"""
from copy import copy
from pathlib import Path
import json

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
RESULT = ROOT / "824" / "f911abeb-6c72-431a-9e42-33b7fb50a7ab_1787560072928_success.jsonl"
TARGET = ROOT / "8.19" / "病例分析题.xlsx"
REQUIREMENTS = ROOT / "data" / "requirements" / "急诊与麻醉科缺失.xlsx"

HEADERS = [
    "题型", "来源行", "专业基地", "科室", "病种", "公共题干", "小问序号", "小问新增信息",
    "小问", "选项", "正确答案", "关键答案", "选项角色", "解析", "难度", "认知层次",
    "考核要点（要求文件）", "质检状态", "质检提示",
]
MEDICAL_FLAGS = {
    "case-妇产-黄体破裂-trial03": "医学复核：月经周期与黄体期时间线需确认。",
    "case-神经内-呼吸衰竭-trial01": "医学复核：吉兰-巴雷综合征治疗选项需确认。",
    "case-眼-心肺复苏-trial05": "医学复核：恢复自主循环后血管活性药用法需确认。",
    "case-耳鼻咽喉-急性会厌炎-trial06": "医学复核：悬雍垂偏斜体征需确认。",
}


def short_base(base):
    return base.replace("基地", "").replace("科", "")


def load_requirement_map():
    ws = openpyxl.load_workbook(REQUIREMENTS, read_only=True, data_only=True)["Sheet1"]
    mapping = {}
    for base, department, disease, _, _, _, case_count in list(ws.values)[1:]:
        if case_count:
            mapping[(short_base(base), disease)] = (base, department, disease)
    return mapping


def resolve_requirement(custom_id, mapping):
    parts = custom_id.rsplit("-", 1)[0].split("-")
    if len(parts) < 3 or parts[0] != "case":
        raise ValueError(f"无法识别 custom_id: {custom_id}")
    key = (parts[1], "-".join(parts[2:]))
    if key not in mapping:
        raise ValueError(f"未匹配需求文件: {custom_id}")
    return mapping[key]


def parse_case(content):
    value = json.loads(content)
    if isinstance(value, list) and len(value) == 1:
        return value[0], []
    if isinstance(value, dict) and isinstance(value.get("case"), dict):
        return value["case"], []
    raise ValueError("不是单案例 JSON")


def common_stem(case, notes):
    values = [(name, case.get(name)) for name in ("common_stem", "common_st", "common_stm") if isinstance(case.get(name), str)]
    if not values:
        raise ValueError("缺少公共题干")
    canonical = case.get("common_stem")
    if not isinstance(canonical, str):
        canonical = values[0][1]
        notes.append(f"字段已规范化：{values[0][0]}→common_stem")
    elif len(values) > 1 and any(value != canonical for _, value in values):
        notes.append("重复公共题干字段内容不一致")
    return canonical.strip()


def validate_case(case):
    notes = []
    stem = common_stem(case, notes)
    questions = case.get("questions")
    if not isinstance(questions, list) or len(questions) != 5:
        raise ValueError("小问数量不是5")
    for index, question in enumerate(questions, 1):
        if not isinstance(question, dict) or question.get("question_no") != index:
            raise ValueError(f"第{index}问序号不正确")
        options = question.get("options")
        if not isinstance(options, list) or not 6 <= len(options) <= 8:
            raise ValueError(f"第{index}问选项数不在6～8")
        labels = [option.get("label") for option in options if isinstance(option, dict)]
        if labels != [chr(65 + item) for item in range(len(options))]:
            raise ValueError(f"第{index}问选项标签不连续")
        answer_labels = [option["label"] for option in options if option.get("role") in {"关键选项", "正确选项"}]
        key_labels = [option["label"] for option in options if option.get("role") == "关键选项"]
        if question.get("answers") != answer_labels or question.get("key_answers") != key_labels:
            raise ValueError(f"第{index}问答案与选项角色不一致")
        if not all(isinstance(question.get(key), str) for key in ("new_information", "stem", "explanation")):
            raise ValueError(f"第{index}问缺少文本字段")
        if len(question["explanation"].strip()) < 80:
            notes.append(f"第{index}问解析较短")
    if sum(len(question["answers"]) >= 2 for question in questions) < 2:
        raise ValueError("多选小问不足2个")
    if not any(question["key_answers"] for question in questions):
        raise ValueError("没有关键选项")
    return stem, notes


def format_options(options):
    return "\n".join(f"{item['label']}. {item['text']}" for item in options)


def format_roles(options):
    return "\n".join(f"{item['label']}={item['role']}" for item in options)


def main():
    requirements = load_requirement_map()
    responses = [json.loads(line) for line in RESULT.read_text(encoding="utf-8").splitlines() if line]
    prepared = []
    for line_number, response in enumerate(responses, 1):
        if response.get("response", {}).get("status_code") != 200:
            raise ValueError(f"第{line_number}行 API 状态异常")
        choice = response["response"]["body"]["choices"][0]
        if choice.get("finish_reason") != "stop":
            raise ValueError(f"第{line_number}行输出未正常结束")
        case, _ = parse_case(choice["message"]["content"])
        stem, notes = validate_case(case)
        base, department, disease = resolve_requirement(response["custom_id"], requirements)
        notes.append("属性已按需求文件回填")
        if response["custom_id"] in MEDICAL_FLAGS:
            notes.append(MEDICAL_FLAGS[response["custom_id"]])
        status = "待人工医学复核" if len(notes) > 1 else "通过基本结构检查"
        prepared.append((line_number, response["custom_id"], base, department, disease, stem, case, status, "；".join(notes)))

    wb = openpyxl.load_workbook(TARGET)
    ws = wb.active
    # 用户已授权清空旧病例分析题；只保留工作表并重建为新题型字段。
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for line_number, _, base, department, disease, stem, case, status, notes in prepared:
        source = f"824/{RESULT.name}:{line_number}"
        for question in case["questions"]:
            ws.append([
                "案例分析题（病例串型不定项选择题）", source, base, department, disease, stem,
                question["question_no"], question["new_information"], question["stem"],
                format_options(question["options"]), "、".join(question["answers"]),
                "、".join(question["key_answers"]), format_roles(question["options"]),
                question["explanation"], case.get("difficulty", ""), case.get("cognitive_level", ""),
                disease, status, notes,
            ])

    widths = [30, 54, 18, 30, 35, 70, 12, 42, 45, 52, 16, 16, 38, 60, 12, 16, 35, 20, 58]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(TARGET)
    print(f"已导入案例 {len(prepared)} 个，小问记录 {len(prepared) * 5} 行：{TARGET}")


if __name__ == "__main__":
    main()
