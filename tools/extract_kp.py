#!/usr/bin/env python3
"""
从2025执医A2真题中反推知识点。
读取原xlsx，在副本上增加知识点列，同时生成去重的知识点表。
"""

import re
import openpyxl
from collections import OrderedDict

# 系统分类关键词映射
SYSTEM_KEYWORDS = {
    "精神科": ["焦虑", "抑郁", "精神", "心理", "强迫", "恐惧", "躁狂", "幻觉", "妄想", "失眠", "人格", "应激", "癔症", "创伤后"],
    "神经内科": ["脑卒中", "癫痫", "帕金森", "偏头痛", "脑梗", "脑出血", "蛛网膜", "脊髓", "周围神经", "重症肌无力", "格林巴利"],
    "呼吸内科": ["肺炎", "肺脓肿", "哮喘", "慢阻肺", "COPD", "肺癌", "胸腔积液", "气胸", "肺栓塞", "肺结核", "呼吸衰竭", "咳嗽", "咯血", "呼吸困难", "支气管"],
    "心内科": ["心肌", "冠心病", "高血压", "心律失常", "心力衰竭", "心包", "瓣膜", "房颤", "室壁瘤", "心内膜", "主动脉", "心脏", "心绞痛", "心梗"],
    "消化内科": ["胃炎", "溃疡", "肝硬化", "胰腺", "胆囊", "肠炎", "结肠炎", "克罗恩", "消化道出血", "食管", "腹痛", "腹泻", "黄疸", "阑尾", "肝癌", "胃癌"],
    "肾内科": ["肾炎", "肾病综合征", "肾衰竭", "透析", "肾小球", "肾盂", "蛋白尿", "血尿", "尿毒症"],
    "血液内科": ["白血病", "淋巴瘤", "贫血", "血小板", "骨髓", "凝血", "溶血", "再障"],
    "内分泌科": ["糖尿病", "甲亢", "甲减", "甲状腺", "肾上腺", "垂体", "痛风", "骨质疏松", "低血糖"],
    "风湿免疫科": ["类风湿", "红斑狼疮", "干燥综合征", "强直", "血管炎", "风湿", "免疫"],
    "普外科": ["疝", "肠梗阻", "结直肠", "乳腺", "甲状腺手术", "胆石症"],
    "骨科": ["骨折", "脱位", "骨囊肿", "骨肿瘤", "椎间盘", "脊柱", "关节", "半月板", "韧带", "股骨", "肱骨", "胫骨"],
    "泌尿外科": ["前列腺", "膀胱", "肾结石", "尿路", "尿道", "睾丸", "阴囊", "肾积水", "泌尿"],
    "胸外科": ["食管癌", "纵隔", "胸壁"],
    "神经外科": ["颅内", "脑膜瘤", "脑外伤", "硬膜", "脑疝"],
    "妇产科": ["妊娠", "分娩", "子宫", "卵巢", "宫颈", "阴道", "月经", "不孕", "流产", "胎盘", "产程", "胎儿", "胎位"],
    "儿科": ["新生儿", "小儿", "婴儿", "早产", "川崎", "先天性"],
    "眼科": ["白内障", "青光眼", "视网膜", "眼底", "屈光", "角膜", "葡萄膜"],
    "耳鼻喉科": ["鼻窦", "中耳", "扁桃体", "喉", "耳聋", "眩晕", "鼻出血"],
    "皮肤科": ["湿疹", "银屑病", "荨麻疹", "皮炎", "带状疱疹", "真菌", "丹毒"],
    "急诊医学": ["休克", "中毒", "创伤", "心肺复苏"],
    "预防医学": ["流行病", "统计", "公共卫生", "职业病", "营养", "霍乱"],
    "医学伦理": ["伦理", "知情同意", "医患", "保密"],
    "卫生法规": ["法规", "执业", "医疗事故", "传染病防治", "药品管理"],
}

# 常见疾病名（用于从选项中识别）
DISEASE_SUFFIXES = ["炎", "癌", "瘤", "症", "病", "综合征", "障碍", "中毒", "结石", "疝", "裂", "疮"]


def extract_diagnosis(explanation: str, answer_label: str, options: list) -> str:
    """从解析和选项中提取知识点名称。"""
    if not explanation:
        return ""

    # 方法1: 从解析中找疾病诊断名 — 最准确
    diag_patterns = [
        r"最可能的诊断是(.{2,12}?)[（(，。]",
        r"首先考虑的诊断是(.{2,12}?)[（(，。]",
        r"考虑为(.{2,12}?)[（(，。]",
        r"诊断为(.{2,12}?)[（(，。]",
        r"最可能的(.{2,8}?)是(.{2,10}?)[（(，。]",
    ]
    for pattern in diag_patterns:
        match = re.search(pattern, explanation)
        if match:
            groups = match.groups()
            diag = groups[-1].strip()  # 取最后一个捕获组
            diag = re.sub(r"^[是为：:]+", "", diag)
            if is_valid_diagnosis(diag):
                return diag

    # 方法2: 从"X对"前面提取疾病名
    # 模式：",XXX（A对）" 或 "。XXX（A对）"
    pattern = r"[，。]([^，。]{2,12}?)（" + re.escape(answer_label) + "对）"
    match = re.search(pattern, explanation)
    if match:
        diag = match.group(1).strip()
        if is_valid_diagnosis(diag):
            return diag

    # 方法3: 从正确答案选项中取
    opt_idx = ord(answer_label) - ord('A')
    if 0 <= opt_idx < len(options):
        opt = options[opt_idx].strip()
        if opt and is_valid_knowledge_point(opt):
            return opt

    return ""


def is_valid_diagnosis(name: str) -> bool:
    """判断是否是合理的疾病诊断名称。"""
    if not name or len(name) < 2 or len(name) > 12:
        return False
    # 排除前缀异常
    if name[0] in "的该和与是为在从到对于把被让可直接由于这":
        return False
    # 排除含标点
    if re.search(r"[，。、：；！？""''（）%+]", name):
        return False
    # 排除句子碎片
    if len(name) > 6 and "是" in name:
        return False
    # 排除非疾病内容
    bad_words = [
        "患者", "考虑", "检查", "治疗", "措施", "原因", "处理", "劝导",
        "健康", "权利", "人格", "事件", "改善", "补充", "体重", "头围",
        "胸围", "囟门", "解释", "直接", "按摩", "引流", "切开",
        "营养", "全身", "情况",
    ]
    for w in bad_words:
        if w in name:
            return False
    # 排除纯数字
    if re.match(r"^[\d\s.%]+$", name):
        return False
    if name[0].isdigit():
        return False
    # 排除药物名
    drug_suffixes = ["素", "醇", "胺", "啶", "苷"]
    if len(name) <= 4 and any(name.endswith(s) for s in drug_suffixes):
        return False
    return True


def is_valid_knowledge_point(name: str) -> bool:
    """判断是否是有效的知识点名称（比诊断更宽松，包括治疗、检查等）。"""
    if not name or len(name) < 2 or len(name) > 12:
        return False
    if name[0] in "的该和与是为在从到对于把被让由于这":
        return False
    if re.search(r"[，。、：；！？""''（）%+]", name):
        return False
    if len(name) > 6 and "是" in name:
        return False
    # 排除非医学内容
    bad_words = ["患者", "事件", "劝导", "健康", "权利", "人格", "解释", "直接",
                 "体重", "头围", "胸围", "囟门", "身长", "营养", "全身", "情况",
                 "改善", "补充", "称为", "符合", "体现", "表明"]
    for w in bad_words:
        if w in name:
            return False
    if re.match(r"^[\d\s.%]+$", name):
        return False
    if name[0].isdigit():
        return False
    # 排除纯测量值
    if re.match(r".*\d+(kg|cm|mm|ml|g|次|分|岁|月|天)", name):
        return False
    return True


def classify_system(stem: str, diagnosis: str, explanation: str) -> str:
    """根据题干和诊断判断所属系统。"""
    text = (stem + diagnosis + explanation)[:500]

    scores = {}
    for system, keywords in SYSTEM_KEYWORDS.items():
        score = sum(1 for kw in keywords if kw in text)
        if score > 0:
            scores[system] = score

    if scores:
        return max(scores, key=scores.get)
    return "其他"


def extract_keywords(stem: str, diagnosis: str, options: list) -> list:
    """提取关键词。"""
    keywords = []

    # 诊断是最重要的关键词
    if diagnosis:
        keywords.append(diagnosis)

    # 从选项中提取疾病名
    for opt in options:
        if opt and 2 <= len(opt) <= 12 and is_valid_diagnosis(opt):
            keywords.append(opt)

    # 从题干提取症状
    symptom_patterns = [
        r"(胸痛|腹痛|头痛|腰痛|关节痛|背痛)",
        r"(发热|咳嗽|咳痰|咯血|呼吸困难|气短)",
        r"(恶心|呕吐|腹泻|便秘|黄疸|呕血|便血)",
        r"(心悸|头晕|乏力|水肿|晕厥|意识模糊)",
        r"(尿频|尿急|尿痛|血尿|蛋白尿|少尿)",
        r"(皮疹|瘙痒|压痛)",
    ]
    for pattern in symptom_patterns:
        matches = re.findall(pattern, stem)
        keywords.extend(matches)

    # 去重
    seen = set()
    result = []
    for kw in keywords:
        if kw not in seen and kw:
            seen.add(kw)
            result.append(kw)
    return result[:6]


def main():
    print("读取原数据集...")
    wb = openpyxl.load_workbook("2025执医a2.xlsx")
    ws = wb.active

    # 收集所有题目数据
    questions = []
    for row_idx in range(2, ws.max_row + 1):
        row_data = [ws.cell(row=row_idx, column=c).value for c in range(1, 10)]
        if not row_data[1]:
            continue
        questions.append({
            "row": row_idx,
            "id": row_data[0],
            "stem": row_data[1] or "",
            "answer": row_data[2] or "",
            "explanation": row_data[3] or "",
            "opt_a": row_data[4] or "",
            "opt_b": row_data[5] or "",
            "opt_c": row_data[6] or "",
            "opt_d": row_data[7] or "",
            "opt_e": row_data[8] or "",
        })

    print(f"共 {len(questions)} 道题目，开始提取知识点...")

    kp_set = OrderedDict()
    results = []

    for i, q in enumerate(questions):
        options = [q["opt_a"], q["opt_b"], q["opt_c"], q["opt_d"], q["opt_e"]]
        diagnosis = extract_diagnosis(q["explanation"], q["answer"], options)
        system = classify_system(q["stem"], diagnosis, q["explanation"])
        keywords = extract_keywords(q["stem"], diagnosis, options)

        kp_name = diagnosis if diagnosis else "待归类"
        kp_key = f"{system}-{kp_name}"

        if kp_key not in kp_set:
            kp_set[kp_key] = {
                "id": f"KP-{len(kp_set)+1:03d}",
                "subject": "临床医学",
                "system": system,
                "name": kp_name,
                "keywords": keywords,
            }

        kp_info = kp_set[kp_key]
        results.append({
            **q,
            "kp_id": kp_info["id"],
            "kp_subject": "临床医学",
            "kp_system": system,
            "kp_name": kp_name,
            "kp_keywords": "、".join(keywords),
        })

        if (i + 1) % 100 == 0:
            print(f"  已处理 {i+1}/{len(questions)}")

    print(f"提取完成，共 {len(kp_set)} 个知识点")

    # 写入副本xlsx
    print("写入带知识点的副本...")
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "A2真题+知识点"

    headers = ["id", "题干", "正确答案", "解析", "选项A", "选项B", "选项C", "选项D", "选项E",
               "知识点ID", "科目", "系统", "知识点名称", "关键词"]
    for col, h in enumerate(headers, 1):
        ws_out.cell(row=1, column=col, value=h)

    for row_idx, r in enumerate(results, 2):
        for col, key in enumerate(["id", "stem", "answer", "explanation",
                                    "opt_a", "opt_b", "opt_c", "opt_d", "opt_e",
                                    "kp_id", "kp_subject", "kp_system", "kp_name", "kp_keywords"], 1):
            ws_out.cell(row=row_idx, column=col, value=r[key])

    wb_out.save("2025执医a2_含知识点.xlsx")
    print("已保存: 2025执医a2_含知识点.xlsx")

    # 写入知识点表
    print("写入知识点表...")
    wb_kp = openpyxl.Workbook()
    ws_kp = wb_kp.active
    ws_kp.title = "知识点"

    kp_headers = ["知识点ID", "科目", "系统", "知识点名称", "关键词"]
    for col, h in enumerate(kp_headers, 1):
        ws_kp.cell(row=1, column=col, value=h)

    for row_idx, (key, kp) in enumerate(kp_set.items(), 2):
        ws_kp.cell(row=row_idx, column=1, value=kp["id"])
        ws_kp.cell(row=row_idx, column=2, value=kp["subject"])
        ws_kp.cell(row=row_idx, column=3, value=kp["system"])
        ws_kp.cell(row=row_idx, column=4, value=kp["name"])
        ws_kp.cell(row=row_idx, column=5, value="、".join(kp["keywords"]))

    wb_kp.save("知识点表.xlsx")
    print(f"已保存: 知识点表.xlsx ({len(kp_set)} 个知识点)")

    # 统计
    print("\n=== 各系统知识点分布 ===")
    system_counts = {}
    for kp in kp_set.values():
        system_counts[kp["system"]] = system_counts.get(kp["system"], 0) + 1
    for system, count in sorted(system_counts.items(), key=lambda x: -x[1]):
        print(f"  {system}: {count} 个")

    classified = sum(1 for kp in kp_set.values() if kp["name"] != "待归类")
    unclassified = len(kp_set) - classified
    print(f"\n已归类: {classified} | 待归类: {unclassified} | 归类率: {classified/len(kp_set)*100:.1f}%")

    # 打印样例
    print("\n=== 知识点样例 ===")
    shown = 0
    for key, kp in kp_set.items():
        if shown >= 15:
            break
        if kp["name"] != "待归类":
            shown += 1
            print(f"  {kp['id']} | {kp['system']} | {kp['name']} | {'、'.join(kp['keywords'][:4])}")


if __name__ == "__main__":
    main()
