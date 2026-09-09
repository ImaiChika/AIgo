#!/usr/bin/env python3
"""删除 8.19 Excel 末尾标记为待复核的 820 行。"""
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]


def main():
    for name in ("A3.xlsx", "A4.xlsx", "病例分析题.xlsx"):
        path = ROOT / "8.19" / name
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        rows = [
            row
            for row in range(2, ws.max_row + 1)
            if isinstance(ws.cell(row, 1).value, str) and "待复核" in ws.cell(row, 1).value
        ]
        if rows and rows != list(range(rows[0], ws.max_row + 1)):
            raise RuntimeError(f"{name} 的待复核行不连续，已停止删除")
        if rows:
            ws.delete_rows(rows[0], len(rows))
        wb.save(path)
        print(f"{name}: 删除待复核行 {len(rows)} 行")


if __name__ == "__main__":
    main()
