#!/usr/bin/env python3
"""删除病例分析题 Excel 的内部质检列。"""
from pathlib import Path

import openpyxl


PATH = Path(__file__).resolve().parents[1] / "8.19" / "病例分析题.xlsx"


def main():
    wb = openpyxl.load_workbook(PATH)
    ws = wb.active
    headers = [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]
    for header in ("质检提示", "质检状态"):
        if header not in headers:
            raise RuntimeError(f"缺少列：{header}")
        column = headers.index(header) + 1
        ws.delete_cols(column)
        headers.pop(column - 1)
    wb.save(PATH)
    print(f"已删除内部质检列：{PATH}")


if __name__ == "__main__":
    main()
