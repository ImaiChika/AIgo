#!/usr/bin/env python3
"""把 820 批次中自动拒绝的原始题组追加到 Excel 末尾，供人工复核。"""
from pathlib import Path
import json
import re

import openpyxl
from openpyxl.styles import PatternFill

from append_819_results import extract_items, valid_group
from append_820_results import NEW_FILES, NEW_DIR, RESULT_DIR


REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")
ILLEGAL_CHARS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def text(value):
    return ILLEGAL_CHARS.sub("", value).strip() if isinstance(value, str) else ""


def options_text(question):
    options = question.get("options") if isinstance(question, dict) else None
    if isinstance(options, list):
        values = []
        for index, option in enumerate(options):
            label = chr(ord("A") + index)
            if isinstance(option, dict):
                label = text(option.get("label")) or label
                value = text(option.get("text"))
            else:
                value = text(option)
            values.append(f"{label}. {value}" if value else label)
        return "\n".join(values)
    return text(options)


def review_reason(question_type, groups, group):
    reasons = []
    if len(groups) != 1:
        reasons.append(f"单请求返回 {len(groups)} 个题组")
    if not isinstance(group, dict):
        reasons.append("题组不是 JSON 对象")
    elif not text(group.get("clinical_stem")):
        reasons.append("病例描述为空")
    elif not valid_group(group, question_type):
        reasons.append("未通过自动结构/长度质检")
    return "；".join(reasons) or "未通过自动质检"


def append_raw_group(ws, question_type, source_ref, groups, group):
    reason = review_reason(question_type, groups, group)
    if not isinstance(group, dict):
        ws.append([f"{question_type}（待复核）", source_ref, "", json.dumps(group, ensure_ascii=False), "", "", "", "", "", reason])
        return 1

    clinical_stem = text(group.get("clinical_stem"))
    questions = group.get("questions")
    if not isinstance(questions, list) or not questions:
        ws.append(
            [
                f"{question_type}（待复核）",
                source_ref,
                clinical_stem or text(group.get("question")),
                "",
                options_text(group),
                text(group.get("answer")),
                text(group.get("explanation", group.get("analysis"))),
                text(group.get("difficulty", "")),
                text(group.get("cognitive_level")),
                f"【待复核原因】{reason}",
            ]
        )
        return 1

    rows = 0
    for question in questions:
        question = question if isinstance(question, dict) else {}
        points = text(question.get("exam_points")) or text(group.get("exam_points"))
        ws.append(
            [
                f"{question_type}（待复核）",
                source_ref,
                clinical_stem,
                text(question.get("stem")),
                options_text(question),
                text(question.get("answer")),
                text(question.get("explanation", question.get("analysis"))),
                text(question.get("difficulty", group.get("difficulty", ""))),
                text(question.get("cognitive_level")) or text(group.get("cognitive_level")),
                f"【待复核原因】{reason}" + (f"；{points}" if points else ""),
            ]
        )
        rows += 1
    return rows


def process(question_type):
    result_path = NEW_DIR / NEW_FILES[question_type]
    responses = [json.loads(line) for line in result_path.read_text(encoding="utf-8").splitlines() if line]
    xlsx_path = RESULT_DIR / f"{question_type}.xlsx"
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    existing_sources = {row[1] for row in ws.iter_rows(min_row=2, values_only=True) if isinstance(row[1], str)}
    appended_groups = appended_rows = 0

    for line_number, response in enumerate(responses, 1):
        source_ref = f"820/{NEW_FILES[question_type]}:{line_number} [待复核]"
        if source_ref in existing_sources:
            continue
        choices = response.get("response", {}).get("body", {}).get("choices", [])
        content = choices[0].get("message", {}).get("content", "") if choices else ""
        groups = extract_items(content)
        rejected = len(groups) != 1 or not groups or not valid_group(groups[0], question_type)
        if not rejected:
            continue
        if not groups:
            ws.append([f"{question_type}（待复核）", source_ref, "", "", "", "", "", "", "", "【待复核原因】未解析出题组"])
            appended_groups += 1
            appended_rows += 1
            continue
        for group in groups:
            appended_rows += append_raw_group(ws, question_type, source_ref, groups, group)
            appended_groups += 1

    start = ws.max_row - appended_rows + 1
    for row in ws.iter_rows(min_row=max(2, start), max_row=ws.max_row):
        for cell in row:
            cell.fill = REVIEW_FILL
    wb.save(xlsx_path)
    return appended_groups, appended_rows


def main():
    for question_type in NEW_FILES:
        groups, rows = process(question_type)
        print(f"{question_type}: 追加待复核题组 {groups} 个，记录 {rows} 行")


if __name__ == "__main__":
    main()
